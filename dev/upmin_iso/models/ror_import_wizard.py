import base64
import io
import re
from datetime import datetime

from odoo import models, fields, api
from odoo.exceptions import ValidationError

import openpyxl

_COPY_FIELDS = [
    "interested_parties",
    "needs_and_exp",
    "compliance",
    "risk_existing_control",
    "opportunities_existing_control",
]

_MAX_ROW_SCAN = 500
_HEADER_ROW_SCAN = 10

_YEAR_RE = re.compile(r"^\d{4}$")
_REVIEW_DATE_RE = re.compile(r"Review Date:?\s*(.+)", re.IGNORECASE)
_SEPT_RE = re.compile(r"\bsept\b", re.IGNORECASE)
_DATE_FORMATS = ("%B %d, %Y", "%b %d, %Y", "%B %d,%Y", "%b %d,%Y")

_RISKS_COL_KEYWORDS = ["inherent", "residual"]
_REQUIRED_ACTION_KEYWORDS = ["required action"]
_RESPONSIBLE_KEYWORDS = ["responsible"]
_STATUS_KEYWORDS = ["status"]


def _cell_text(value):
    return str(value).strip() if value not in (None, "") else ""


def _norm(text):
    return re.sub(r"\s+", " ", text).strip().lower()


def _contains_any(text, keywords):
    norm = _norm(text)
    return any(kw in norm for kw in keywords)


def _parse_flexible_date(text):
    text = _cell_text(text)
    if not text:
        return None
    normalized = _SEPT_RE.sub("Sep", text)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(normalized, fmt).date()
        except ValueError:
            continue
    return None


def _split_rich_pair(value, label_a, label_b):
    """Split 'Label A: text A' text back into text_a, tolerating either label prefix."""
    text = _cell_text(value)
    for label in (label_a, label_b):
        prefix = f"{label}:"
        if _norm(text).startswith(_norm(prefix)):
            return text[len(prefix):].strip()
    return text


def _to_selection_value(value):
    """Normalize a L/F/S rating cell (may be int, float, or str) to '1'-'4', else False."""
    if isinstance(value, (int, float)):
        as_int = int(value)
        if as_int == value and 1 <= as_int <= 4:
            return str(as_int)
        return False
    text = _cell_text(value)
    return text if text in {"1", "2", "3", "4"} else False


def _split_responsible_date(text):
    if not text or text == "-":
        return False, False
    responsible, _, date_str = text.partition("/")
    responsible = responsible.strip() or False
    due_date = _parse_flexible_date(date_str.strip())
    return responsible, due_date


class RORImportWizard(models.TransientModel):
    _name = "upmin_iso.ror_import_wizard"
    _description = "Import ROR from XLSX"

    ror_id = fields.Many2one(
        comodel_name="upmin_iso.ror",
        string="ROR",
        required=True,
    )
    import_file = fields.Binary("Import File", required=True)
    import_filename = fields.Char("File Name")
    result_summary = fields.Text("Result", readonly=True)

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        active_id = self.env.context.get("active_id")
        if active_id and "ror_id" in fields_list:
            res["ror_id"] = active_id
        return res

    def action_import_ror(self):
        self.ensure_one()
        if not self.import_file:
            raise ValidationError("Please upload a file to import.")

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.import_file)), data_only=True
            )
        except Exception as exc:
            raise ValidationError(f"Could not read the uploaded file as an XLSX workbook: {exc}")

        stats = {"issues_created": 0, "issues_updated": 0, "ratings_created": 0, "ratings_updated": 0}
        warnings = []
        for sheet in workbook.worksheets:
            if not _YEAR_RE.match(str(sheet.title).strip()):
                continue  # not a ROR year sheet (e.g. a "Criteria" reference tab)
            self._import_sheet(sheet, stats, warnings)

        summary = (
            f"Issues: {stats['issues_created']} created, {stats['issues_updated']} updated.\n"
            f"Ratings: {stats['ratings_created']} created, {stats['ratings_updated']} updated."
        )
        if warnings:
            summary += "\n\nWarnings:\n" + "\n".join(f"- {w}" for w in warnings)
        self.result_summary = summary

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def _find_header_row(self, sheet):
        max_row = min(sheet.max_row, _HEADER_ROW_SCAN)
        max_col = sheet.max_column
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                if _contains_any(_cell_text(sheet.cell(row=r, column=c).value), ["requirement/issue"]):
                    return r
        return None

    def _find_quarter_blocks(self, sheet, header_row):
        """Locate quarter blocks by scanning header_row for 'Review Date' cells.

        Each match's column is the block's end column. A block's start column is the
        previous block's end + 1, or the first column (after B-H) whose header mentions
        inherent/residual risks if it's the first block found.
        """
        max_col = sheet.max_column
        review_cols = []
        for c in range(9, max_col + 1):
            text = _cell_text(sheet.cell(row=header_row, column=c).value)
            if _REVIEW_DATE_RE.search(text):
                review_cols.append(c)

        blocks = []
        prev_end = None
        for end_col in review_cols:
            if prev_end is None:
                start_col = None
                for c in range(9, end_col + 1):
                    text = _cell_text(sheet.cell(row=header_row, column=c).value)
                    if _contains_any(text, _RISKS_COL_KEYWORDS):
                        start_col = c
                        break
                if start_col is None:
                    start_col = 9
            else:
                start_col = prev_end + 1
            blocks.append((start_col, end_col))
            prev_end = end_col
        return blocks

    def _find_col_by_keyword(self, sheet, row, start_col, end_col, keywords):
        for c in range(start_col, end_col + 1):
            text = _cell_text(sheet.cell(row=row, column=c).value)
            if _contains_any(text, keywords):
                return c
        return None

    def _find_lfs_cols(self, sheet, header_row, start_col, end_col):
        """Locate the L / F / S single-letter marker columns within a block."""
        lfs = {}
        for r in (header_row, header_row + 1, header_row + 2):
            for c in range(start_col, end_col + 1):
                text = _cell_text(sheet.cell(row=r, column=c).value).upper()
                if text in ("L", "F", "S") and text not in lfs:
                    lfs[text] = c
        return lfs.get("L"), lfs.get("F"), lfs.get("S")

    def _import_sheet(self, sheet, stats, warnings):
        header_row = self._find_header_row(sheet)
        if header_row is None:
            warnings.append(f"Sheet '{sheet.title}': couldn't locate a 'Requirement/Issue' header row — skipped.")
            return

        quarter_blocks = self._find_quarter_blocks(sheet, header_row)
        if not quarter_blocks:
            warnings.append(f"Sheet '{sheet.title}': no 'Review Date' columns found — issues will be imported without ratings.")

        for start_col, end_col in quarter_blocks:
            header_text = _cell_text(sheet.cell(row=header_row, column=end_col).value)
            match = _REVIEW_DATE_RE.search(header_text)
            parsed = _parse_flexible_date(match.group(1)) if match else None
            if parsed and str(parsed.year) != str(sheet.title).strip():
                warnings.append(
                    f"Sheet '{sheet.title}': quarter header references {parsed.year} "
                    f"('{match.group(1).strip()}') — please verify this date."
                )

        section_col = 2
        row = header_row + 1
        max_row = min(sheet.max_row, _MAX_ROW_SCAN)

        current_type = None
        found_section = False
        while row <= max_row:
            label = _cell_text(sheet.cell(row=row, column=section_col).value)
            if _norm(label) == "internal issues":
                current_type = "W"
                found_section = True
                row += 1
                continue
            if _norm(label) == "external issues":
                current_type = "T"
                found_section = True
                row += 1
                continue
            if not label:
                row += 1
                continue

            if current_type is None:
                row += 1
                continue

            row_number = _cell_text(sheet.cell(row=row, column=1).value)
            if not row_number:
                # Not a numbered issue row (e.g. trailing legend/notes text) — current
                # section has ended; keep scanning in case another section follows.
                current_type = None
                row += 1
                continue

            self._import_issue_block(sheet, row, current_type, quarter_blocks, header_row, stats, warnings)
            row += 2

        if not found_section:
            warnings.append(f"Sheet '{sheet.title}': no 'Internal Issues' or 'External Issues' section found — skipped.")

    def _import_issue_block(self, sheet, row, swot_type, quarter_blocks, header_row, stats, warnings):
        ror = self.ror_id

        def cell(r, c):
            return sheet.cell(row=r, column=c).value

        vals = {
            "description": _cell_text(cell(row, 2)),
            "interested_parties": _cell_text(cell(row, 3)),
            "needs_and_exp": _cell_text(cell(row, 4)),
            "compliance": _cell_text(cell(row, 5)),
            "risks": _split_rich_pair(cell(row, 6), "Risk", "Opportunity"),
            "opportunities": _split_rich_pair(cell(row + 1, 6), "Opportunity", "Risk"),
            "consequence": _split_rich_pair(cell(row, 7), "Consequence", "Benefit"),
            "benefit": _split_rich_pair(cell(row + 1, 7), "Benefit", "Consequence"),
            "risk_existing_control": _cell_text(cell(row, 8)),
            "opportunities_existing_control": _cell_text(cell(row + 1, 8)),
            "swot_type": swot_type,
            "ror_id": ror.id,
        }

        if not vals["description"]:
            warnings.append(f"Sheet '{sheet.title}', row {row}: issue is missing its description — skipped.")
            return

        issues = ror.internal_issues if swot_type == "W" else ror.external_issues
        issue = issues.filtered(lambda i: i.description == vals["description"])[:1]
        if issue:
            issue.write({k: v for k, v in vals.items() if k not in ("swot_type", "ror_id")})
            stats["issues_updated"] += 1
        else:
            issue = self.env["upmin_iso.swot_line"].create(vals)
            stats["issues_created"] += 1

        for start_col, end_col in quarter_blocks:
            self._import_rating_block(sheet, row, start_col, end_col, header_row, issue, stats, warnings)

    def _import_rating_block(self, sheet, row, start_col, end_col, header_row, issue, stats, warnings):
        header_text = _cell_text(sheet.cell(row=header_row, column=end_col).value)
        match = _REVIEW_DATE_RE.search(header_text)
        if not match:
            return

        review_date = _parse_flexible_date(match.group(1))
        if not review_date:
            warnings.append(
                f"Sheet '{sheet.title}', column {end_col}: could not parse review date "
                f"'{match.group(1).strip()}' — quarter skipped."
            )
            return

        l_col, f_col, s_col = self._find_lfs_cols(sheet, header_row, start_col, end_col)
        action_col = self._find_col_by_keyword(sheet, header_row, start_col, end_col, _REQUIRED_ACTION_KEYWORDS)
        resp_col = self._find_col_by_keyword(sheet, header_row, start_col, end_col, _RESPONSIBLE_KEYWORDS)

        def sel(col):
            if not col:
                return False
            return _to_selection_value(sheet.cell(row=row, column=col).value)

        def sel2(col):
            if not col:
                return False
            return _to_selection_value(sheet.cell(row=row + 1, column=col).value)

        if not any([sel(l_col), sel(f_col), sel(s_col), sel2(l_col), sel2(f_col), sel2(s_col)]):
            # Placeholder quarter for this issue ("-"/blank cells) — no rating to import.
            return

        risk_action = _cell_text(sheet.cell(row=row, column=action_col).value) if action_col else ""
        opp_action = _cell_text(sheet.cell(row=row + 1, column=action_col).value) if action_col else ""
        risk_resp_text = _cell_text(sheet.cell(row=row, column=resp_col).value) if resp_col else ""
        opp_resp_text = _cell_text(sheet.cell(row=row + 1, column=resp_col).value) if resp_col else ""
        risk_responsible, risk_due_date = _split_responsible_date(risk_resp_text)
        opp_responsible, opp_due_date = _split_responsible_date(opp_resp_text)
        risk_status = _cell_text(sheet.cell(row=row, column=end_col).value) if end_col != resp_col else ""
        opp_status = _cell_text(sheet.cell(row=row + 1, column=end_col).value) if end_col != resp_col else ""

        vals = {
            "risk_likelihood": sel(l_col),
            "opportunity_likelihood": sel2(l_col),
            "risk_frequency": sel(f_col),
            "opportunity_frequency": sel2(f_col),
            "consequence_severity": sel(s_col),
            "benefit_severity": sel2(s_col),
            "risk_required_action": risk_action or False,
            "opportunity_required_action": opp_action or False,
            "risk_responsible": risk_responsible or False,
            "opportunity_responsible": opp_responsible or False,
            "risk_due_date": risk_due_date,
            "opportunity_due_date": opp_due_date,
            "risk_status": risk_status or False,
            "opportunity_status": opp_status or False,
        }

        RORRating = self.env["upmin_iso.ror_rating"]
        rating = RORRating.search([("issue", "=", issue.id), ("review_date", "=", review_date)], limit=1)
        try:
            with self.env.cr.savepoint():
                if rating:
                    rating.write(vals)
                else:
                    vals.update({"issue": issue.id, "review_date": review_date})
                    RORRating.create(vals)
        except ValidationError as exc:
            warnings.append(
                f"Sheet '{sheet.title}', column {end_col}: rating for '{review_date}' "
                f"rejected — {exc.args[0] if exc.args else exc} — quarter skipped."
            )
            return

        if rating:
            stats["ratings_updated"] += 1
        else:
            stats["ratings_created"] += 1
